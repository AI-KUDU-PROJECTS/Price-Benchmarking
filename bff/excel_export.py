"""Create the Marketing Excel export from the same BFF contract as React."""
from __future__ import annotations

import sqlite3
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from adapters.base import BrandAdapter
from bff.contract import Product, now_riyadh

HEADERS = [
    "Brand",
    "Channel",
    "Category",
    "Name (English)",
    "Name (Arabic)",
    "Regular Price",
    "Special Price",
    "Effective Price",
    "Currency",
    "Availability",
    "Status",
    "Location",
    "Observed At",
    "Source Run ID",
    "Product ID",
    "Image URL",
]


def _row(product: Product) -> list[object | None]:
    effective_price = (
        product.special_price
        if product.special_price is not None
        else product.regular_price
    )
    return [
        product.brand_id,
        product.channel,
        product.category,
        product.name_en,
        product.name_ar,
        product.regular_price,
        product.special_price,
        effective_price,
        product.currency,
        product.availability,
        product.status,
        product.location,
        product.observed_at,
        product.source_run_id,
        product.id,
        product.image_url,
    ]


def build_market_workbook(adapters: dict[str, BrandAdapter]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Current Prices"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:P1"

    header_fill = PatternFill("solid", fgColor="1A458A")
    for index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    products: list[Product] = []
    for adapter in adapters.values():
        try:
            products.extend(adapter.list_products())
        except (FileNotFoundError, sqlite3.Error):
            continue
    products.sort(
        key=lambda item: (
            item.brand_id,
            item.channel,
            item.category or "",
            item.name_en or item.name_ar or "",
        )
    )
    for product in products:
        sheet.append(_row(product))

    widths = [16, 16, 24, 34, 34, 15, 15, 15, 12, 14, 18, 28, 24, 34, 42, 54]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    metadata = workbook.create_sheet("Metadata")
    metadata.append(["Generated At", now_riyadh().isoformat(timespec="seconds")])
    metadata.append(["Product Rows", len(products)])
    metadata.append(["Source", "KUDU Price Benchmark BFF"])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
