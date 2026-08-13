"""
tests/test_excel_exporter.py
Covers required test #13: Excel generation.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from competitors.herfy.backend import config, excel_exporter
from competitors.herfy.tests.test_change_detector import bump_run, ingest


def test_daily_report_generates_all_required_worksheets(conn, tmp_path, monkeypatch, delivery_result, pickup_result):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path / "exports")
    ingest(conn, tmp_path, "batch1", "DELIVERY", bump_run(delivery_result, "2026-08-01"))
    ingest(conn, tmp_path, "batch1", "PICKUP", bump_run(pickup_result, "2026-08-01"))

    path = excel_exporter.export_daily_report(conn, 29696, date(2026, 8, 1))
    assert path.exists()
    assert path.name.startswith("Herfy_Daily_Changes_")
    wb = load_workbook(path)
    expected_sheets = {
        "Summary", "Pickup Products", "Pickup Offers", "Delivery Products",
        "Delivery Offers", "Daily Changes", "Not Observed", "Sources and Run Metadata",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))
    assert wb["Delivery Products"].max_row > 1
    assert wb["Pickup Products"].max_row > 1
    # The "Offers" CMS category (see backend/offer_parser.py) DOES populate
    # for this brand - the fixture has 2 (Double Offer, Offer you'll love)
    # per channel, so header + 2 data rows.
    assert wb["Pickup Offers"].max_row == 3
    assert wb["Delivery Offers"].max_row == 3


def test_monthly_report_generates_all_required_worksheets(conn, tmp_path, monkeypatch, delivery_result):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path / "exports")
    ingest(conn, tmp_path, "batch1", "DELIVERY", bump_run(delivery_result, "2026-07-15"))
    ingest(conn, tmp_path, "batch2", "DELIVERY", bump_run(delivery_result, "2026-08-15"))

    path = excel_exporter.export_monthly_report(conn, 29696, 2026, 8)
    assert path.exists()
    assert path.name.startswith("Herfy_Monthly_Comparison_")
    wb = load_workbook(path)
    expected_sheets = {
        "Monthly Summary", "Pickup Comparison", "Delivery Comparison", "New Products",
        "New Offers", "Ended Offers", "Price Changes", "Removed and Returned", "Sources and Run Metadata",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))


def test_current_catalog_generates_all_required_worksheets(conn, tmp_path, monkeypatch, delivery_result, pickup_result):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path / "exports")
    ingest(conn, tmp_path, "batch1", "DELIVERY", bump_run(delivery_result, "2026-08-01"))
    ingest(conn, tmp_path, "batch1", "PICKUP", bump_run(pickup_result, "2026-08-01"))

    path = excel_exporter.export_current_catalog(conn, 29696, date(2026, 8, 1))
    assert path.exists()
    assert path.name.startswith("Herfy_Current_Catalog_")
    wb = load_workbook(path)
    expected_sheets = {"Pickup Normalized", "Pickup Legacy View", "Delivery Normalized", "Delivery Legacy View", "Offers"}
    assert expected_sheets.issubset(set(wb.sheetnames))

    # Legacy View must show the REAL per-size price for "Beef Tortilla
    # Meal" (Regular=29, Medium=32, Large=34 - see normalizer.py's
    # extract_sizes()), and leave every size column blank for a product
    # with no confirmed "Sizes" option group (e.g. the Jalapeno Bites).
    legacy_ws = wb["Delivery Legacy View"]
    header = [c.value for c in legacy_ws[1]]
    assert header == ["Category", "Item", "Sandwich", "Regular", "Medium", "Large", "Source Link"]
    rows_by_item = {row[1]: row for row in legacy_ws.iter_rows(min_row=2, values_only=True)}
    assert rows_by_item["Beef Tortilla Meal"][3:6] == (29, 32, 34)
    assert rows_by_item["4 pcs. Jalapeno Bites with Cheese"][3:6] == (None, None, None)

    # Regression: the Normalized sheet's "Sizes" column must render each
    # size as "Title: price" text, never Python's raw dict repr (e.g.
    # "{'title': 'Regular', 'price': 29}") - a real bug found live where
    # _join() called bare str() on each size dict.
    normalized_ws = wb["Delivery Normalized"]
    norm_header = [c.value for c in normalized_ws[1]]
    norm_rows = {row[norm_header.index("Product")]: row for row in normalized_ws.iter_rows(min_row=2, values_only=True)}
    sizes_cell = norm_rows["Beef Tortilla Meal"][norm_header.index("Sizes")]
    assert sizes_cell == "Regular: 29, Medium: 32, Large: 34"
    assert "{" not in sizes_cell and "'" not in sizes_cell
