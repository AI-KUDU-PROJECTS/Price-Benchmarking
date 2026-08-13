"""
tests/test_excel_exporter.py
Covers required test #13: Excel generation.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from competitors.burger_king.backend import config, excel_exporter
from competitors.burger_king.tests.test_change_detector import bump_run, ingest


def test_daily_report_generates_all_required_worksheets(conn, tmp_path, monkeypatch, delivery_result, pickup_result):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path / "exports")
    ingest(conn, tmp_path, "batch1", "DELIVERY", bump_run(delivery_result, "2026-08-01"))
    ingest(conn, tmp_path, "batch1", "PICKUP", bump_run(pickup_result, "2026-08-01"))

    path = excel_exporter.export_daily_report(conn, 11474, date(2026, 8, 1))
    assert path.exists()
    assert path.name.startswith("BurgerKing_Daily_Changes_")
    wb = load_workbook(path)
    expected_sheets = {
        "Summary", "Pickup Products", "Pickup Offers", "Delivery Products",
        "Delivery Offers", "Daily Changes", "Not Observed", "Sources and Run Metadata",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))
    assert wb["Delivery Products"].max_row > 1
    assert wb["Pickup Products"].max_row > 1
    # KING DAILY DEALS products (see backend/offer_parser.py) DO populate the
    # Offers sheets - the fixture has 2 (Cheeseburger Lovers, King Wrap Box)
    # per channel, so header + 2 data rows.
    assert wb["Pickup Offers"].max_row == 3
    assert wb["Delivery Offers"].max_row == 3


def test_monthly_report_generates_all_required_worksheets(conn, tmp_path, monkeypatch, delivery_result):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path / "exports")
    ingest(conn, tmp_path, "batch1", "DELIVERY", bump_run(delivery_result, "2026-07-15"))
    ingest(conn, tmp_path, "batch2", "DELIVERY", bump_run(delivery_result, "2026-08-15"))

    path = excel_exporter.export_monthly_report(conn, 11474, 2026, 8)
    assert path.exists()
    assert path.name.startswith("BurgerKing_Monthly_Comparison_")
    wb = load_workbook(path)
    expected_sheets = {
        "Monthly Summary", "Pickup Comparison", "Delivery Comparison", "New Products",
        "New Offers", "Ended Offers", "Price Changes", "Removed and Returned", "Sources and Run Metadata",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))


def test_current_catalog_generates_all_required_worksheets(conn, tmp_path, monkeypatch, delivery_result, pickup_result):
    monkeypatch.setattr(config, "EXPORTS_DIR", tmp_path / "exports")
    ingest(conn, tmp_path, "batch1", "DELIVERY", bump_run(delivery_result, "2026-08-01"))
    ingest(conn, tmp_path, "batch1", "PICKUP", bump_run(pickup_result, "2026-08-01"))

    path = excel_exporter.export_current_catalog(conn, 11474, date(2026, 8, 1))
    assert path.exists()
    assert path.name.startswith("BurgerKing_Current_Catalog_")
    wb = load_workbook(path)
    expected_sheets = {"Pickup Normalized", "Pickup Legacy View", "Delivery Normalized", "Delivery Legacy View", "Offers"}
    assert expected_sheets.issubset(set(wb.sheetnames))

    # Legacy View must never fabricate a size - none of the fixture
    # products carry a recognized size variant, so every Regular/Medium/
    # Large column must be blank, not a guess.
    legacy_ws = wb["Delivery Legacy View"]
    header = [c.value for c in legacy_ws[1]]
    assert header == ["Category", "Item", "Sandwich", "Regular", "Medium", "Large", "Source Link"]
    for row in legacy_ws.iter_rows(min_row=2, values_only=True):
        assert row[3] is None and row[4] is None and row[5] is None
